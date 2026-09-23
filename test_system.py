"""
Verification test suite for Apple Stock vs iPhone calculations and consistency.
"""

import os
import openpyxl
from update_sheet import IPHONE_DATASET, fetch_live_aapl_price

def test_dataset_completeness():
    print(f"Checking dataset completeness: {len(IPHONE_DATASET)} models loaded.")
    assert len(IPHONE_DATASET) == 53, f"Expected 53 models, got {len(IPHONE_DATASET)}"
    print("Dataset completeness: PASS")

def test_mathematical_consistency():
    price = 325.70
    total_cost = sum(m[2] for m in IPHONE_DATASET)
    total_shares = sum(m[2] / m[3] for m in IPHONE_DATASET)
    total_value = total_shares * price
    total_profit = total_value - total_cost
    overall_roi = total_profit / total_cost

    print(f"Total iPhone Cost: ${total_cost:,.2f}")
    print(f"Total AAPL Shares: {total_shares:.2f}")
    print(f"Total Invested Value @ ${price}: ${total_value:,.2f}")
    print(f"Total Net Profit: ${total_profit:,.2f}")
    print(f"Overall Cumulative ROI: {overall_roi*100:.2f}%")

    # Verify original iPhone (2007) invariants
    iphone_2007 = IPHONE_DATASET[0]
    assert iphone_2007[0] == "iPhone"
    shares_2007 = iphone_2007[2] / iphone_2007[3] # 499 / 4.36 = 114.4495
    val_2007 = shares_2007 * price
    roi_2007 = (val_2007 / iphone_2007[2]) - 1
    assert round(shares_2007, 2) == 114.45
    print(f"iPhone (2007) ROI: +{roi_2007*100:.2f}% (Shares: {shares_2007:.4f})")
    print("Mathematical Invariants: PASS")

def test_workbook_integrity():
    wb_path = os.path.join(os.path.dirname(__file__), "outsoucrd apple sheet.xlsx")
    assert os.path.exists(wb_path), f"Workbook missing at {wb_path}"
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    ws = wb.active
    assert ws.max_row == 55, f"Expected 55 rows (headers + 53 models + 1 summary), got {ws.max_row}"
    # Verify row 2 formulas
    assert ws.cell(2, 6).value == "=C2/D2*E2"
    assert ws.cell(2, 7).value == "=F2-C2"
    assert ws.cell(2, 8).value == "=(F2/C2)-1"
    print("Workbook Integrity & Formulas: PASS")

if __name__ == "__main__":
    test_dataset_completeness()
    test_mathematical_consistency()
    test_workbook_integrity()
    print("ALL TESTS PASSED SUCCESSFULLY!")
